"""Builders for the tiny documents the extraction tests need.

Generating the fixtures instead of committing binaries keeps the repository honest: a test that
reads a PDF says how that PDF was made, and a reviewer can change the page count without a hex dump.
"""

from __future__ import annotations

import io
from collections.abc import Sequence

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

PAGE_WIDTH = 595
PAGE_HEIGHT = 842


def digital_pdf(*pages: Sequence[str]) -> bytes:
    """A PDF with a real text layer, one page per argument, one line per string."""
    return _pdf([_text_stream(lines) for lines in pages])


def blank_pdf(page_count: int = 1) -> bytes:
    """Pages with no text layer at all: the shape of a scan without the scanner."""
    return _pdf([" "] * page_count)


def deck(*slides: Sequence[str]) -> bytes:
    """A real PPTX, one text box per line, top to bottom in reading order."""
    presentation = Presentation()
    layout = presentation.slide_layouts[6]
    for lines in slides:
        slide = presentation.slides.add_slide(layout)
        top = Inches(1)
        for line in lines:
            box = slide.shapes.add_textbox(Inches(1), top, Inches(6), Inches(1))
            box.text_frame.text = line
            top = top + Inches(1.2)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def deck_with_picture(png: bytes) -> bytes:
    """A deck whose only content is an image, the case that has to reach the recognizer."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1), Inches(4), Inches(2))
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def white_png(width: int = 400, height: int = 200) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, format="PNG")
    return buffer.getvalue()


def _text_stream(lines: Sequence[str]) -> str:
    drawn = " ".join(f"({_escape(line)}) Tj 0 -16 Td" for line in lines if line)
    return f"BT /F1 12 Tf 72 760 Td {drawn} ET"


def _escape(line: str) -> str:
    return line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _pdf(contents: list[str]) -> bytes:
    """Minimal PDF with a valid cross-reference table: no fonts beyond Helvetica."""
    page_count = len(contents)
    page_ids = list(range(4, 4 + page_count))
    content_ids = list(range(4 + page_count, 4 + 2 * page_count))
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    objects.extend(
        (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {PAGE_WIDTH} {PAGE_HEIGHT}] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        for content_id in content_ids
    )
    for content in contents:
        payload = content.encode("latin-1", errors="replace")
        header = b"<< /Length " + str(len(payload)).encode() + b" >>\nstream\n"
        objects.append(header + payload + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    out += trailer.encode()
    return bytes(out)
